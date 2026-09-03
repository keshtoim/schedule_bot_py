from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

_WS = str | Path

_WHITESPACE_RE = re.compile(r"\s+")


@dataclass
class Sheet:
    """Thin wrapper over an openpyxl worksheet that transparently resolves
    merged cells and exposes the bits the parsers need."""

    _ws: Worksheet
    _merge_anchor: dict[tuple[int, int], tuple[int, int]] = field(repr=False)
    merges: list[str]

    @property
    def max_row(self) -> int:
        return self._ws.max_row

    @property
    def max_column(self) -> int:
        return self._ws.max_column

    def text(self, row: int, col: int) -> str:
        """Cell text with whitespace collapsed. A cell inside a merged range
        resolves to the range's top-left ("anchor") value, like ExcelJS."""
        anchor = self._merge_anchor.get((row, col), (row, col))
        value = self._ws.cell(row=anchor[0], column=anchor[1]).value
        if value is None:
            return ""
        return _WHITESPACE_RE.sub(" ", str(value)).strip()


def load_active_sheet(file_path: _WS) -> Sheet:
    """The workbook has one visible worksheet (the current data) plus a stack
    of hidden historical versions — no date parsing of sheet names needed."""
    workbook = load_workbook(filename=str(file_path), data_only=True, read_only=False)
    active = next((ws for ws in workbook.worksheets if ws.sheet_state == "visible"), None)
    if active is None:
        raise ValueError(f"No visible worksheet found in {file_path}")

    merge_anchor: dict[tuple[int, int], tuple[int, int]] = {}
    merges: list[str] = []
    for rng in active.merged_cells.ranges:
        merges.append(str(rng.coord))
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                merge_anchor[(row, col)] = (rng.min_row, rng.min_col)

    return Sheet(_ws=active, _merge_anchor=merge_anchor, merges=merges)
